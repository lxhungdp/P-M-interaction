"""Full calculation workbook export (Excel) — inputs, outputs, mesh, P-M chart."""

from __future__ import annotations

import io
import math
from datetime import datetime
from typing import Any, Dict, List, Optional

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from src.analysis.solver import pm_results_to_dataframe, position_summary_to_dataframe
from src.materials.concrete import Concrete
from src.materials.rebar import Rebar
from src.materials.strand import Strand

from pm_app.config import MESH_AG_TOL, MESH_CB_TOL, MESH_NY_CAP
from pm_app.models import AnalysisInputs
from pm_app.pm_checks import is_safe, m_cap, ray_intersect_pm_curve, utilization


def _fiber_rows(section, theta_rad: float) -> pd.DataFrame:
    cos_t = math.cos(theta_rad)
    sin_t = math.sin(theta_rad)
    rows = []
    for i, f in enumerate(section.fibers, start=1):
        if isinstance(f.material, Concrete):
            mat = "concrete"
        elif isinstance(f.material, Rebar):
            mat = "rebar"
        elif isinstance(f.material, Strand):
            mat = "strand"
        else:
            mat = type(f.material).__name__
        yp = -f.x * sin_t + f.y * cos_t
        rows.append({
            "fiber_id": i,
            "material": mat,
            "x_mm": round(f.x, 4),
            "y_mm": round(f.y, 4),
            "y_prime_mm": round(yp, 4),
            "area_mm2": round(f.area, 4),
        })
    return pd.DataFrame(rows)


def _inputs_df(inputs: AnalysisInputs) -> pd.DataFrame:
    rows = [
        ("code_name", inputs.code_name),
        ("fck [MPa]", inputs.fck),
        ("fy [MPa]", inputs.fy),
        ("curve_conc", inputs.curve_conc),
        ("sec_type", inputs.sec_type),
        ("b [mm]", inputs.b),
        ("h [mm]", inputs.h),
        ("cover [mm]", inputs.cover),
        ("bar_dia [mm]", inputs.bar_dia),
        ("n_top", inputs.n_top),
        ("n_bot", inputs.n_bot),
        ("n_left", inputs.n_left),
        ("n_right", inputs.n_right),
        ("use_strand", inputs.use_strand),
        ("strand_dia_val [mm]", inputs.strand_dia_val),
        ("fpk_val [MPa]", inputs.fpk_val),
        ("fp01k_val [MPa]", inputs.fp01k_val),
        ("Ep_val [MPa]", inputs.Ep_val),
        ("fpe_val [MPa]", inputs.fpe_val),
        ("eps_pe_val", inputs.eps_pe_val),
        ("avg_Mx [kN·m]", inputs.avg_Mx),
        ("avg_My [kN·m]", inputs.avg_My),
        ("theta_rad", inputs.theta_rad),
        ("theta_deg_auto", inputs.theta_deg_auto),
        ("sym_check", inputs.sym_check),
        ("n_pts", inputs.n_pts),
        ("ny_fiber (input)", inputs.ny_fiber),
    ]
    return pd.DataFrame(rows, columns=["parameter", "value"])


def _rebar_input_df(inputs: AnalysisInputs) -> pd.DataFrame:
    rows: List[dict] = []
    for i, item in enumerate(inputs.rb_list or [], start=1):
        if len(item) >= 3:
            x, y, dia = item[0], item[1], item[2]
        elif len(item) >= 2:
            x, y, dia = item[0], item[1], inputs.bar_dia
        else:
            continue
        rows.append({"no": i, "x_mm": x, "y_mm": y, "D_mm": dia, "source": "rb_list"})
    for i, pos in enumerate(inputs.poly_rebar_positions or [], start=1):
        if len(pos) >= 2:
            rows.append({
                "no": len(rows) + 1,
                "x_mm": pos[0],
                "y_mm": pos[1],
                "D_mm": inputs.bar_dia,
                "source": "poly_rebar",
            })
    return pd.DataFrame(rows)


def _strand_input_df(inputs: AnalysisInputs) -> pd.DataFrame:
    if not inputs.use_strand:
        return pd.DataFrame(columns=["no", "x_mm", "y_mm", "D_mm"])
    rows = [
        {"no": i + 1, "x_mm": p[0], "y_mm": p[1], "D_mm": inputs.strand_dia_val}
        for i, p in enumerate(inputs.strand_pos or [])
    ]
    return pd.DataFrame(rows)


def _material_df(cache: dict, code: Any = None) -> pd.DataFrame:
    rows = [
        {"item": "code", "value": cache.get("code_name", "")},
        {"item": "fcd [MPa]", "value": cache["fcd_val"]},
        {"item": "fyd [MPa]", "value": cache["fyd_val"]},
        {"item": "εcu", "value": cache["ecu_val"]},
        {"item": "εcu [‰]", "value": cache["ecu_val"] * 1e3},
        {"item": "Ec [MPa]", "value": cache["Ec_val"]},
        {"item": "Es [MPa]", "value": 200000.0},
        {"item": "MESH_AG_TOL [%]", "value": MESH_AG_TOL * 100},
        {"item": "MESH_CB_TOL [%]", "value": MESH_CB_TOL * 100},
        {"item": "MESH_NY_CAP", "value": MESH_NY_CAP},
    ]
    if code is not None:
        try:
            mf = code.material_factors("ULS")
            rows.extend([
                {"item": "gamma_c (Φc)", "value": mf.gamma_c},
                {"item": "gamma_s (Φs)", "value": mf.gamma_s},
                {"item": "alpha_cc", "value": mf.alpha_cc},
            ])
        except Exception:
            pass
    return pd.DataFrame(rows)


def _geometry_df(cache: dict) -> pd.DataFrame:
    geo = cache["geo"]
    section = cache["section"]
    a_st = sum(f.area for f in section.fibers if isinstance(f.material, (Rebar, Strand)))
    a_c = sum(f.area for f in section.fibers if isinstance(f.material, Concrete))
    rows = [
        {"property": "Ag_theory [mm²]", "value": cache["Ag_theory"]},
        {"property": "Ag_mesh gross [mm²]", "value": section.gross_area()},
        {"property": "Ac_concrete [mm²]", "value": a_c},
        {"property": "As_steel [mm²]", "value": a_st},
    ]
    try:
        for k, v in geo.to_dict().items():
            rows.append({"property": k, "value": v})
    except Exception:
        rows.extend([
            {"property": "geo.A [mm²]", "value": geo.A},
            {"property": "cx [mm]", "value": geo.cx},
            {"property": "cy [mm]", "value": geo.cy},
            {"property": "Ixx [mm⁴]", "value": geo.Ixx},
            {"property": "Iyy [mm⁴]", "value": geo.Iyy},
            {"property": "Iy_prime [mm⁴]", "value": geo.Iy_prime},
            {"property": "y_prime_top [mm]", "value": geo.y_prime_top},
            {"property": "y_prime_bot [mm]", "value": geo.y_prime_bot},
        ])
    return pd.DataFrame(rows)


def _output_results_df(cache: dict) -> pd.DataFrame:
    solver_out = cache["solver_out"]
    n_nom = np.asarray(cache["N_nom"], float)
    rows = [
        ("N_max_design [kN]", cache["N_max_v"]),
        ("N_max_nominal [kN]", float(n_nom.max())),
        ("N_bal [kN]", cache["N_bal"]),
        ("M_bal [kN·m]", cache["M_bal"]),
        ("N_p0 [kN]", cache["N_p0"]),
        ("M_p0 [kN·m]", cache["M_p0"]),
        ("c_bal [mm]", cache["c_bal_mm"]),
        ("eps_bal_top", cache["eps_bal_top"]),
        ("eps_bal_bot", cache["eps_bal_bot"]),
        ("eps_bal_top [‰]", cache["eps_bal_top"] * 1e3),
        ("eps_bal_bot [‰]", cache["eps_bal_bot"] * 1e3),
        ("lc1_Pn_cap [kN]", cache.get("lc1_Pn_cap")),
        ("lc1_Mn_cap [kN·m]", cache.get("lc1_Mn_cap")),
        ("lc1_eps_bot", cache.get("lc1_eps_bot")),
        ("h_eff_solver [mm]", getattr(solver_out, "h_eff_mm", None)),
        ("n_pts", cache.get("n_pts")),
        ("ny_fiber_req", cache.get("ny_fiber_req")),
        ("ny_fiber_used", cache.get("ny_fiber_used")),
        ("mesh_warn", cache.get("mesh_warn") or ""),
        ("Ag_mesh_normalized", cache.get("Ag_mesh_normalized", False)),
        ("concrete_area_scale", cache.get("ag_mesh_scale", 1.0)),
        ("n_fibers_total", len(cache["section"].fibers)),
        ("theta_deg", cache.get("theta_deg_auto")),
    ]
    return pd.DataFrame(rows, columns=["result", "value"])


def _key_points_df(cache: dict) -> pd.DataFrame:
    return pd.DataFrame([
        {"point": "Pmax (pure compression)", "M_kNm": 0.0, "N_kN": cache["N_max_v"]},
        {"point": "Balanced (Mb, Pb)", "M_kNm": cache["M_bal"], "N_kN": cache["N_bal"]},
        {"point": "P≈0 (M0)", "M_kNm": cache["M_p0"], "N_kN": cache["N_p0"]},
        {"point": "LC1 ray capacity", "M_kNm": cache.get("lc1_Mn_cap"), "N_kN": cache.get("lc1_Pn_cap")},
    ])


def _pm_curve_df(cache: dict) -> pd.DataFrame:
    n_des = cache["N_des"]
    m_des = cache["M_des"]
    n_nom = cache["N_nom"]
    m_nom = cache["M_nom"]
    return pd.DataFrame({
        "point_no": range(1, len(n_des) + 1),
        "M_design_kNm": m_des,
        "N_design_kN": n_des,
        "M_nominal_kNm": m_nom,
        "N_nominal_kN": n_nom,
    })


def _pm_load_points_df(cache: dict, sym_check: bool) -> pd.DataFrame:
    n_des = np.asarray(cache["N_des"], float)
    m_des = np.asarray(cache["M_des"], float)
    rows: List[dict] = []
    for lc in cache.get("load_cases") or []:
        try:
            pu = float(lc["Pu [kN]"])
            mx = float(lc["Mx [kN·m]"])
            my = float(lc["My [kN·m]"])
            mu = math.sqrt(mx ** 2 + my ** 2)
            mc = m_cap(pu, n_des, m_des)
            util = utilization(pu, mu, n_des, m_des)
            rows.append({
                "case": lc.get("케이스", "LC"),
                "extreme": bool(lc.get("극단")),
                "Pu_kN": pu,
                "Mx_kNm": mx,
                "My_kNm": my,
                "Mu_kNm": mu,
                "M_capacity_kNm": mc,
                "UR": util,
                "OK": is_safe(pu, mu, n_des, m_des),
                "plot_M_kNm": mu,
                "plot_N_kN": pu,
            })
            if sym_check:
                rows.append({
                    "case": f"{lc.get('케이스', 'LC')} (sym -M)",
                    "extreme": bool(lc.get("극단")),
                    "Pu_kN": pu,
                    "Mx_kNm": mx,
                    "My_kNm": my,
                    "Mu_kNm": -mu,
                    "M_capacity_kNm": mc,
                    "UR": util,
                    "OK": is_safe(pu, mu, n_des, m_des),
                    "plot_M_kNm": -mu,
                    "plot_N_kN": pu,
                })
        except Exception:
            continue
    if rows and cache.get("lc1_Mn_cap") is not None:
        rows.append({
            "case": "LC1_capacity",
            "extreme": False,
            "Pu_kN": None,
            "Mx_kNm": None,
            "My_kNm": None,
            "Mu_kNm": cache.get("lc1_Mn_cap"),
            "M_capacity_kNm": None,
            "UR": None,
            "OK": None,
            "plot_M_kNm": cache.get("lc1_Mn_cap"),
            "plot_N_kN": cache.get("lc1_Pn_cap"),
        })
    return pd.DataFrame(rows)


def _load_check_df(cache: dict) -> pd.DataFrame:
    n_des = np.asarray(cache["N_des"], float)
    m_des = np.asarray(cache["M_des"], float)
    rows = []
    for lc in cache.get("load_cases") or []:
        try:
            pu = float(lc["Pu [kN]"])
            mx = float(lc["Mx [kN·m]"])
            my = float(lc["My [kN·m]"])
            mu = math.sqrt(mx ** 2 + my ** 2)
            mc = m_cap(pu, n_des, m_des)
            util = utilization(pu, mu, n_des, m_des)
            rows.append({
                "case": lc.get("케이스", "LC"),
                "extreme": "★" if lc.get("극단") else "",
                "Pu_kN": round(pu, 1),
                "Mx_kNm": round(mx, 1),
                "My_kNm": round(my, 1),
                "Mu_kNm": round(mu, 1),
                "M_capacity_kNm": round(mc, 1),
                "UR": f"{util:.1%}",
                "OK": "OK" if is_safe(pu, mu, n_des, m_des) else "NG",
            })
        except Exception:
            continue
    return pd.DataFrame(rows)


def _summary_df(cache: dict, inputs: AnalysisInputs) -> pd.DataFrame:
    return pd.DataFrame([
        {"item": "exported_at", "value": datetime.now().isoformat(timespec="seconds")},
        {"item": "title", "value": "단면 P-M상관도 검토 — full calculation export"},
        {"item": "sec_type", "value": inputs.sec_type},
        {"item": "section_b×h [mm]", "value": f"{inputs.b:.0f}×{inputs.h:.0f}"},
        {"item": "code_name", "value": inputs.code_name},
        {"item": "mesh_warn", "value": cache.get("mesh_warn") or ""},
        {"item": "ny_input", "value": cache.get("ny_fiber_req")},
        {"item": "ny_used", "value": cache.get("ny_fiber_used")},
        {"item": "Ag_normalized", "value": cache.get("Ag_mesh_normalized", False)},
        {"item": "concrete_area_scale", "value": cache.get("ag_mesh_scale", 1.0)},
        {"item": "n_mesh_iterations", "value": len(cache.get("mesh_iterations") or [])},
        {"item": "n_slice_rows", "value": len(cache.get("mesh_slice_rows") or [])},
        {"item": "n_fibers_final", "value": len(cache["section"].fibers)},
        {"item": "PM_chart_sheet", "value": "PM_Chart (tab 1) — image + Excel chart"},
    ])


def _code_summary_df(code: Any) -> pd.DataFrame:
    if code is None:
        return pd.DataFrame(columns=["code_summary"])
    try:
        text = code.summary()
    except Exception:
        text = ""
    return pd.DataFrame([{"line_no": i + 1, "text": line} for i, line in enumerate(text.splitlines())])


def _render_pm_chart_png(cache: dict, sym_check: bool) -> bytes:
    """Matplotlib P-M diagram (M horizontal, N vertical) matching app tab."""
    m_des = np.asarray(cache["M_des"], float)
    n_des = np.asarray(cache["N_des"], float)
    m_nom = np.asarray(cache["M_nom"], float)
    n_nom = np.asarray(cache["N_nom"], float)

    fig, ax = plt.subplots(figsize=(8, 6))
    ax.plot(m_nom, n_nom, "--", color="#90A4AE", linewidth=1.8, label="Nominal (φ=1)")
    ax.plot(m_des, n_des, "-", color="#1565C0", linewidth=2.2, label="Design")
    ax.fill_between(m_des, 0, n_des, alpha=0.06, color="#1565C0")

    ax.scatter([0], [cache["N_max_v"]], marker="*", s=120, color="#C62828", zorder=5, label="Pmax")
    ax.scatter([cache["M_bal"]], [cache["N_bal"]], marker="D", s=80, color="#2E7D32", zorder=5, label="Balanced")
    ax.scatter([cache["M_p0"]], [cache["N_p0"]], marker="o", s=60, color="#E65100", zorder=5, label="P≈0")

    load_cases = cache.get("load_cases") or []
    if load_cases:
        try:
            lc0 = load_cases[0]
            n_lc = float(lc0["Pu [kN]"])
            m_lc = math.sqrt(float(lc0["Mx [kN·m]"]) ** 2 + float(lc0["My [kN·m]"]) ** 2)
            mc, nc = ray_intersect_pm_curve(m_lc, n_lc, m_des, n_des)
            if mc is not None:
                ax.plot([0, mc], [0, nc], ":", color="#B71C1C", linewidth=1.5)
                ax.scatter([mc], [nc], marker="x", s=80, color="#B71C1C", zorder=5)
        except Exception:
            pass

    for lc in load_cases:
        try:
            pu = float(lc["Pu [kN]"])
            mu = math.sqrt(float(lc["Mx [kN·m]"]) ** 2 + float(lc["My [kN·m]"]) ** 2)
            signs = [1, -1] if sym_check else [1]
            for s in signs:
                ax.scatter([s * mu], [pu], marker="^", s=50, color="#E53935", zorder=5)
        except Exception:
            pass

    ax.set_xlabel("M [kN·m]")
    ax.set_ylabel("N [kN]")
    ax.set_title(
        f"P-M Interaction  |  b×h = {cache.get('b', 0):.0f}×{cache.get('h', 0):.0f} mm  "
        f"theta = {cache.get('theta_deg_auto', 0):.1f} deg"
    )
    ax.grid(True, alpha=0.3)
    ax.legend(loc="best", fontsize=8)
    ax.set_xlim(left=-max(float(m_des.max()), 1.0) * 0.05 if sym_check else 0)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=120)
    plt.close(fig)
    buf.seek(0)
    return buf.getvalue()


def _add_pm_excel_chart(ws, n_rows: int, anchor: str = "H2") -> None:
    """Native Excel XY line chart (M on X, N on Y)."""
    from openpyxl.chart import Reference, ScatterChart, Series

    if n_rows < 2:
        return
    chart = ScatterChart()
    chart.scatterStyle = "lineMarker"
    chart.title = "P-M Interaction (Excel chart)"
    chart.x_axis.title = "M [kN·m]"
    chart.y_axis.title = "N [kN]"
    chart.style = 2
    chart.width = 20
    chart.height = 14

    # PM_Curve columns: point_no, M_design, N_design, M_nominal, N_nominal
    x_des = Reference(ws, min_col=2, min_row=2, max_row=n_rows + 1)
    y_des = Reference(ws, min_col=3, min_row=2, max_row=n_rows + 1)
    chart.series.append(Series(y_des, x_des, title="Design"))

    x_nom = Reference(ws, min_col=4, min_row=2, max_row=n_rows + 1)
    y_nom = Reference(ws, min_col=5, min_row=2, max_row=n_rows + 1)
    chart.series.append(Series(y_nom, x_nom, title="Nominal"))

    ws.add_chart(chart, anchor)


def _setup_pm_chart_sheet(
    pm_ws,
    pm_png: bytes,
    pm_curve_ws,
    n_curve_rows: int,
) -> None:
    """First tab: embedded PNG + Excel chart (high visibility)."""
    from openpyxl.drawing.image import Image as XLImage

    pm_ws["A1"] = "P-M INTERACTION DIAGRAM"
    pm_ws["A2"] = (
        "Image below matches app tab 1. Numeric data: sheets PM_Curve, PM_LoadPoints."
    )
    pm_ws.column_dimensions["A"].width = 100
    for row in range(3, 32):
        pm_ws.row_dimensions[row].height = 18
    pm_ws.row_dimensions[3].height = 290

    img = XLImage(io.BytesIO(pm_png))
    img.width = 760
    img.height = 520
    pm_ws.add_image(img, "A3")

    if n_curve_rows >= 2:
        _add_pm_excel_chart(pm_curve_ws, n_curve_rows, anchor="H2")
        # Duplicate chart on PM_Chart referencing PM_Curve data (same workbook)
        from openpyxl.chart import Reference, ScatterChart, Series

        chart = ScatterChart()
        chart.scatterStyle = "lineMarker"
        chart.title = "P-M (linked to PM_Curve data)"
        chart.x_axis.title = "M [kN·m]"
        chart.y_axis.title = "N [kN]"
        chart.width = 20
        chart.height = 14
        x_des = Reference(pm_curve_ws, min_col=2, min_row=2, max_row=n_curve_rows + 1)
        y_des = Reference(pm_curve_ws, min_col=3, min_row=2, max_row=n_curve_rows + 1)
        chart.series.append(Series(y_des, x_des, title="Design"))
        x_nom = Reference(pm_curve_ws, min_col=4, min_row=2, max_row=n_curve_rows + 1)
        y_nom = Reference(pm_curve_ws, min_col=5, min_row=2, max_row=n_curve_rows + 1)
        chart.series.append(Series(y_nom, x_nom, title="Nominal"))
        pm_ws.add_chart(chart, "L3")


def _autosize_columns(ws, df: pd.DataFrame) -> None:
    from openpyxl.utils import get_column_letter

    for col_idx, col in enumerate(df.columns, start=1):
        lengths = [len(str(col))]
        if len(df) > 0:
            lengths.extend(len(str(v)) for v in df[col].head(80).astype(str))
        max_len = max(lengths) if lengths else 10
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 45)


def build_calc_workbook_bytes(
    cache: dict,
    inputs: AnalysisInputs,
    code: Any = None,
) -> bytes:
    """Build multi-sheet Excel workbook; returns .xlsx bytes."""
    try:
        import openpyxl  # noqa: F401
        from PIL import Image as PILImage  # noqa: F401 — required for embedded PNG
    except ImportError as exc:
        raise ImportError(
            "Excel export requires openpyxl and Pillow. pip install openpyxl Pillow"
        ) from exc

    buf = io.BytesIO()
    theta = float(cache.get("theta_rad", inputs.theta_rad))
    sym_check = bool(cache.get("sym_check", inputs.sym_check))
    solver_out = cache["solver_out"]

    sheets: Dict[str, pd.DataFrame] = {
        "Summary": _summary_df(cache, inputs),
        "Input": _inputs_df(inputs),
        "Input_Rebar": _rebar_input_df(inputs),
        "Input_Strand": _strand_input_df(inputs),
        "LoadCases": inputs.df_loads.copy(),
        "Material": _material_df(cache, code),
        "Code_Summary": _code_summary_df(code),
        "Mesh_Iterations": pd.DataFrame(cache.get("mesh_iterations") or []),
        "Mesh_Slices": pd.DataFrame(cache.get("mesh_slice_rows") or []),
        "Fibers_Final": _fiber_rows(cache["section"], theta),
        "Geometry": _geometry_df(cache),
        "Output_Results": _output_results_df(cache),
        "PM_Curve": _pm_curve_df(cache),
        "PM_LoadPoints": _pm_load_points_df(cache, sym_check),
        "Key_Points": _key_points_df(cache),
        "Solver": pm_results_to_dataframe(solver_out, half_only=True),
        "Position_Strain": position_summary_to_dataframe(solver_out.position_summary),
        "Load_Check": _load_check_df(cache),
    }

    pm_png = _render_pm_chart_png(cache, sym_check)
    n_curve = len(sheets["PM_Curve"])

    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        # PM_Chart first tab so it is visible when opening the file
        pm_ws = writer.book.create_sheet("PM_Chart", 0)

        for name, df in sheets.items():
            safe_name = name[:31]
            df.to_excel(writer, sheet_name=safe_name, index=False)
            _autosize_columns(writer.sheets[safe_name], df)

        _setup_pm_chart_sheet(pm_ws, pm_png, writer.sheets["PM_Curve"], n_curve)

    buf.seek(0)
    return buf.getvalue()
